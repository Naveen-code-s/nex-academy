from pathlib import Path

from langchain_core.documents import Document
from langchain_community.document_loaders import PyMuPDFLoader
from docx import Document as DocxDocument


def load_document(file_path: str):
    """
    Load PDF, DOCX, or TXT files and return LangChain Documents.
    """

    path = Path(file_path)
    extension = path.suffix.lower()

    if extension == ".pdf":
        loader = PyMuPDFLoader(str(path))
        documents = loader.load()

        # Convert page numbers from 0-based to 1-based
        for doc in documents:
            if "page" in doc.metadata:
                doc.metadata["page"] = doc.metadata["page"] + 1

            doc.metadata["source"] = path.name

        return documents

    if extension == ".docx":
        docx = DocxDocument(str(path))

        documents = []

        for index, paragraph in enumerate(docx.paragraphs, start=1):
            text = paragraph.text.strip()

            if text:
                documents.append(
                    Document(
                        page_content=text,
                        metadata={
                            "source": path.name,
                            "page": index
                        }
                    )
                )

        return documents

    if extension == ".txt":
        text = path.read_text(
            encoding="utf-8",
            errors="ignore"
        )

        return [
            Document(
                page_content=text,
                metadata={
                    "source": path.name,
                    "page": 1
                }
            )
        ]

    raise ValueError(
        f"Unsupported file type: {extension}"
    )